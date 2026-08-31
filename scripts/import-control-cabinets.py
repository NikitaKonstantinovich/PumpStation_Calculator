from __future__ import annotations

import hashlib
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ROOT = ROOT / "Equipment" / "наработки для конфигуратора ШУ"
PRICE_ROOT = SOURCE_ROOT / "Прайсы оборудования" / "Прайсы оборудования"
CALCULATION = SOURCE_ROOT / "Расчет стоимости шкафов.xlsm"
PUBLIC_OUTPUT = ROOT / "frontend" / "public" / "control-cabinet-database.json"
SQL_OUTPUT = ROOT / "frontend" / "drizzle" / "0002_control_cabinet_configurator.sql"

SKIPPED_PRICE_FILES = {
    "0_лист с кодом копирующим весь текст сос страницы.xlsm",
    "00_Блок расчета стоимости через чип и дип.xlsm",
}
STATIC_ARTICLES = [
    ("control-relay", "RRP20-3-05-220A", 1),
    ("relay-socket", "RRP20D-RRM-3", 1),
    ("signal-lamp", "25003DEK", 1),
    ("terminal", "scr-ut-4-b", 6),
    ("filter", "DA92", 2),
    ("cross-module", "Б0043931", 1),
    ("cable-gland", "YSA20-14-16-54-K41", 6),
]
ASSEMBLY_KIT_PRICES = {"тип 2.1": 3000, "тип 2.2": 4000, "тип 3.1": 5000, "тип 3.2": 6000}
SMART_POWER_RANGES = [
    (0.37, 0.75, 0.75),
    (0.8, 1.5, 1.5),
    (1.6, 2.2, 2.2),
    (2.4, 4.0, 3.7),
    (4.1, 5.5, 5.5),
    (5.6, 7.5, 7.5),
]


def clean(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        return round(value, 10) if math.isfinite(value) else None
    if isinstance(value, str):
        return value.replace("\u00a0", " ").replace("\r\n", "\n").strip()
    return str(value)


def number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    if not isinstance(value, str):
        return None
    normalized = value.strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    normalized = re.sub(r"[₽рР]+$", "", normalized)
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", normalized):
        return None
    try:
        return float(normalized)
    except ValueError:
        return None


def text_key(value: Any) -> str:
    return str(clean(value) or "").casefold().replace("ё", "е")


def slug(value: str) -> str:
    normalized = value.casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "-", normalized).strip("-") or "item"


def stable_id(*parts: object) -> str:
    digest = hashlib.sha1("|".join(map(str, parts)).encode("utf-8")).hexdigest()[:18]
    return f"cc-{digest}"


def price_files() -> list[Path]:
    files = [*PRICE_ROOT.glob("*.xlsx"), *PRICE_ROOT.glob("*.xlsm")]
    files += [*PRICE_ROOT.glob("*/*.xlsx"), *PRICE_ROOT.glob("*/*.xlsm")]
    return sorted(
        (path for path in files if path.name not in SKIPPED_PRICE_FILES),
        key=lambda path: str(path.relative_to(PRICE_ROOT)).casefold(),
    )


def read_components() -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for source in price_files():
        workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=source.suffix.casefold() == ".xlsm")
        relative = str(source.relative_to(PRICE_ROOT)).replace("\\", "/")
        category = source.stem
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(clean(value) or f"column_{index + 1}") for index, value in enumerate(next(rows, ()))]
            price_indexes = [index for index, header in enumerate(headers) if re.search(r"цен|стоим", header, re.IGNORECASE)]
            price_index = price_indexes[-1] if price_indexes else len(headers) - 1
            for source_row, row in enumerate(rows, start=2):
                values = [clean(value) for value in row]
                article = str(values[0] or "").strip() if values else ""
                manufacturer = str(values[1] or "").strip() if len(values) > 1 else ""
                name = str(values[2] or "").strip() if len(values) > 2 else ""
                component_type = str(values[3] or "").strip() if len(values) > 3 else ""
                if not article and not name:
                    continue
                attributes = {
                    headers[index]: value
                    for index, value in enumerate(values)
                    if index not in {0, 1, 2, 3, price_index} and value not in (None, "")
                }
                current_price = number(values[price_index]) if price_index < len(values) else None
                result.append({
                    "id": stable_id(relative, sheet.title, source_row, article),
                    "category": category,
                    "manufacturer": manufacturer,
                    "article": article,
                    "name": name or component_type or article,
                    "componentType": component_type,
                    "attributes": attributes,
                    "currentPrice": round(current_price, 2) if current_price is not None and current_price > 0 else None,
                    "currency": "RUB",
                    "sourceFile": relative,
                    "sourceSheet": sheet.title,
                    "sourceRow": source_row,
                })
        workbook.close()
    for kit_type, kit_price in ASSEMBLY_KIT_PRICES.items():
        result.append({
            "id": f"synthetic-assembly-{slug(kit_type)}",
            "category": "Сборочный комплект",
            "manufacturer": "Wenard",
            "article": kit_type,
            "name": f"Сборочный комплект {kit_type}",
            "componentType": "Сборочный комплект",
            "attributes": {"Состав": "Провода, кабельный короб, наконечники, маркировка и крепёж"},
            "currentPrice": kit_price,
            "currency": "RUB",
            "sourceFile": "Расчет стоимости шкафов.xlsm",
            "sourceSheet": "Сборочный комплект",
            "sourceRow": 0,
        })
    result.append({
        "id": "synthetic-cabinet-assembly",
        "category": "Работы",
        "manufacturer": "Wenard",
        "article": "LABOR-CABINET-ASSEMBLY",
        "name": "Сборка шкафа управления",
        "componentType": "Работа",
        "attributes": {"Человеко-часы": 10, "Стоимость часа, руб.": 2000},
        "currentPrice": 20000,
        "currency": "RUB",
        "sourceFile": "Расчет стоимости шкафов.xlsm",
        "sourceSheet": "Smart_НС",
        "sourceRow": 0,
    })
    return result


def attr_number(component: dict[str, Any], pattern: str) -> float | None:
    for label, value in component["attributes"].items():
        if re.search(pattern, label, re.IGNORECASE):
            parsed = number(value)
            if parsed is not None:
                return parsed
    return None


def choose(
    components: list[dict[str, Any]],
    *,
    category: str,
    sheet: str | None = None,
    type_contains: str | None = None,
    name_contains: str | None = None,
    numeric_pattern: str | None = None,
    numeric_target: float | None = None,
    exact_numeric: bool = False,
    extra: Any = None,
) -> dict[str, Any]:
    candidates = [
        item for item in components
        if text_key(item["category"]) == text_key(category)
        and (sheet is None or text_key(item["sourceSheet"]) == text_key(sheet))
        and (type_contains is None or text_key(type_contains) in text_key(item["componentType"]))
        and (name_contains is None or text_key(name_contains) in text_key(item["name"]))
        and item["currentPrice"] is not None
        and (extra is None or extra(item))
    ]
    if numeric_pattern is not None and numeric_target is not None:
        rated = [(item, attr_number(item, numeric_pattern)) for item in candidates]
        rated = [(item, rating) for item, rating in rated if rating is not None]
        if exact_numeric:
            rated = [(item, rating) for item, rating in rated if abs(rating - numeric_target) < 1e-7]
        else:
            rated = [(item, rating) for item, rating in rated if rating + 1e-7 >= numeric_target]
        rated.sort(key=lambda pair: (pair[1], pair[0]["currentPrice"]))
        candidates = [item for item, _ in rated]
    if not candidates:
        raise LookupError(f"Component not found: {category=} {sheet=} {type_contains=} {name_contains=} {numeric_target=}")
    return candidates[0]


def by_article(components: list[dict[str, Any]], article: str) -> dict[str, Any]:
    matches = [item for item in components if text_key(item["article"]) == text_key(article) and item["currentPrice"] is not None]
    if not matches:
        raise LookupError(f"Article not found: {article}")
    return sorted(matches, key=lambda item: item["currentPrice"])[0]


def build_bom(components: list[dict[str, Any]], params: dict[str, Any]) -> list[dict[str, Any]]:
    switch = choose(
        components,
        category="Рубильники",
        sheet="dekraft",
        type_contains="ВН-102",
        numeric_pattern=r"номинал,?\s*а",
        numeric_target=params["incomingSwitchCurrentA"],
        exact_numeric=True,
        extra=lambda item: attr_number(item, r"полюс") == 3,
    )
    breaker = choose(
        components,
        category="Автоматический_выключатель",
        sheet="dekraft",
        type_contains="ВА-103",
        numeric_pattern=r"номинал,?\s*а",
        numeric_target=params["breakerCurrentA"],
        exact_numeric=True,
        extra=lambda item: attr_number(item, r"полюс") == 3,
    )
    vfd = choose(
        components,
        category="ПЧ",
        sheet="DA",
        type_contains="DA G3",
        numeric_pattern=r"мощност",
        numeric_target=params["vfdPowerKw"],
    )
    contactor = choose(
        components,
        category="Контактор",
        sheet="dekraft",
        type_contains="КМ-102",
        numeric_pattern=r"мощност",
        numeric_target=params["pumpPowerKw"],
    )
    enclosure = choose(
        components,
        category="Корпуса_шкафов",
        sheet="IEK",
        type_contains="ЩМП",
        extra=lambda item: any(
            text_key(value).replace("x", "х") == text_key(params["enclosureDimensions"]).replace("x", "х")
            for label, value in item["attributes"].items() if re.search(r"габарит", label, re.IGNORECASE)
        ),
    )
    selected = [
        ("dynamic", "incoming-switch", switch, 1),
        ("dynamic", "motor-breaker", breaker, params["pumpCount"]),
        ("dynamic", "vfd", vfd, params["vfdCount"]),
        ("dynamic", "contactor", contactor, params["contactorCount"]),
        ("dynamic", "enclosure", enclosure, 1),
        ("dynamic", "assembly-kit", next(item for item in components if item["id"] == f"synthetic-assembly-{slug(params['assemblyKitType'])}"), 1),
        ("dynamic", "labor", next(item for item in components if item["id"] == "synthetic-cabinet-assembly"), 1),
    ]
    selected += [("static", role, by_article(components, article), quantity) for role, article, quantity in STATIC_ARTICLES]
    return [
        {
            "role": role,
            "componentGroup": component_group,
            "componentId": component["id"],
            "quantity": quantity,
            "sortOrder": index,
            "component": component,
            "currentCost": round((component["currentPrice"] or 0) * quantity, 2),
        }
        for index, (component_group, role, component, quantity) in enumerate(selected, start=1)
        if quantity > 0
    ]


def read_cabinets(components: list[dict[str, Any]]) -> list[dict[str, Any]]:
    workbook = load_workbook(CALCULATION, read_only=True, data_only=True, keep_vba=True)
    sheet = workbook["Smart_НС"]
    cabinets: list[dict[str, Any]] = []
    for row_index in range(3, 19):
        row = [clean(sheet.cell(row_index, column).value) for column in range(1, 14)]
        pump_power = number(row[1])
        if not row[0] or pump_power is None:
            continue
        match = re.search(r"Smart_(\d+)-", str(row[0]), re.IGNORECASE)
        pump_count = int(match.group(1)) if match else 0
        params = {
            "pumpCount": pump_count,
            "pumpPowerKw": pump_power,
            "breakerCurrentA": float(number(row[2]) or 0),
            "incomingSwitchCurrentA": float(number(row[3]) or 0),
            "contactorCount": int(number(row[4]) or 0),
            "vfdCount": int(number(row[5]) or 0),
            "vfdPowerKw": float(number(row[6]) or 0),
            "enclosureDimensions": str(row[7]),
            "assemblyKitType": str(row[8]),
            "laborHours": float(number(row[9]) or 0),
            "laborRate": float(number(row[10]) or 0),
        }
        bom = build_bom(components, params)
        current_total = round(sum(item["currentCost"] for item in bom), 2)
        cabinet_id = f"SMART-{pump_count}-{str(pump_power).replace('.', ',')}"
        cabinets.append({
            "id": cabinet_id,
            "configurationKey": f"smart:{pump_count}:{pump_power:g}",
            "stationType": "smart",
            "name": str(row[0]),
            **params,
            "cachedTotal": round(float(number(row[12]) or current_total), 2),
            "currentTotal": current_total,
            "source": "Расчет стоимости шкафов.xlsm · Smart_НС",
            "items": bom,
        })
    workbook.close()
    grouped: list[dict[str, Any]] = []
    for pump_count in (2, 3):
        for power_min, power_max, representative_power in SMART_POWER_RANGES:
            cabinet = next(
                item for item in cabinets
                if item["pumpCount"] == pump_count
                and abs(item["pumpPowerKw"] - representative_power) < 1e-7
            )
            format_power = lambda value: f"{value:g}".replace(".", ",")
            range_label = f"{format_power(power_min)}-{format_power(power_max)}"
            grouped.append({
                **cabinet,
                "id": f"SMART-{pump_count}-{range_label}",
                "configurationKey": f"smart:{pump_count}:{power_min:g}-{power_max:g}",
                "name": f"Wenard PC BP-Smart-{pump_count}-({range_label})",
                "pumpPowerKw": power_max,
                "powerMinKw": power_min,
                "powerMaxKw": power_max,
                "source": "Расчет стоимости шкафов.xlsm · Smart_НС · Таблица градации.xlsx",
            })
    return grouped


def sql_string(value: Any) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def micros(value: float | None) -> str:
    return "NULL" if value is None else str(round(value * 1_000_000))


def write_sql(components: list[dict[str, Any]], cabinets: list[dict[str, Any]]) -> None:
    lines = [
        'CREATE TABLE `control_components` (',
        '  `id` text PRIMARY KEY NOT NULL,',
        '  `category` text NOT NULL,',
        '  `manufacturer` text,',
        '  `article` text,',
        '  `name` text NOT NULL,',
        '  `component_type` text,',
        '  `attributes_json` text DEFAULT \'{}\' NOT NULL,',
        '  `current_price_microunits` integer,',
        '  `currency` text DEFAULT \'RUB\' NOT NULL,',
        '  `source_file` text NOT NULL,',
        '  `source_sheet` text NOT NULL,',
        '  `source_row` integer NOT NULL,',
        '  `is_active` integer DEFAULT 1 NOT NULL,',
        '  `updated_at` text DEFAULT CURRENT_TIMESTAMP NOT NULL',
        ');',
        '--> statement-breakpoint',
        'CREATE INDEX `idx_control_components_category` ON `control_components` (`category`);',
        '--> statement-breakpoint',
        'CREATE INDEX `idx_control_components_article` ON `control_components` (`article`);',
        '--> statement-breakpoint',
        'CREATE TABLE `control_cabinets` (',
        '  `id` text PRIMARY KEY NOT NULL,',
        '  `configuration_key` text NOT NULL,',
        '  `station_type` text NOT NULL,',
        '  `name` text NOT NULL,',
        '  `pump_count` integer NOT NULL,',
        '  `pump_power_kw` real NOT NULL,',
        '  `breaker_current_a` real NOT NULL,',
        '  `incoming_switch_current_a` real NOT NULL,',
        '  `contactor_count` integer NOT NULL,',
        '  `vfd_count` integer NOT NULL,',
        '  `vfd_power_kw` real NOT NULL,',
        '  `enclosure_dimensions` text NOT NULL,',
        '  `assembly_kit_type` text NOT NULL,',
        '  `labor_hours` real NOT NULL,',
        '  `labor_rate` real NOT NULL,',
        '  `cached_total_microunits` integer NOT NULL,',
        '  `source` text NOT NULL,',
        '  `created_by_user_id` text REFERENCES `users`(`id`) ON DELETE set null,',
        '  `price_updated_at` text DEFAULT CURRENT_TIMESTAMP NOT NULL,',
        '  `created_at` text DEFAULT CURRENT_TIMESTAMP NOT NULL,',
        '  `updated_at` text DEFAULT CURRENT_TIMESTAMP NOT NULL',
        ');',
        '--> statement-breakpoint',
        'CREATE UNIQUE INDEX `uq_control_cabinets_configuration` ON `control_cabinets` (`configuration_key`);',
        '--> statement-breakpoint',
        'CREATE INDEX `idx_control_cabinets_lookup` ON `control_cabinets` (`station_type`,`pump_count`,`pump_power_kw`);',
        '--> statement-breakpoint',
        'CREATE TABLE `control_cabinet_items` (',
        '  `id` integer PRIMARY KEY AUTOINCREMENT NOT NULL,',
        '  `cabinet_id` text NOT NULL REFERENCES `control_cabinets`(`id`) ON DELETE cascade,',
        '  `component_id` text NOT NULL REFERENCES `control_components`(`id`) ON DELETE restrict,',
        '  `role` text NOT NULL,',
        '  `quantity` real NOT NULL,',
        '  `sort_order` integer DEFAULT 0 NOT NULL',
        ');',
        '--> statement-breakpoint',
        'CREATE UNIQUE INDEX `uq_control_cabinet_items_role` ON `control_cabinet_items` (`cabinet_id`,`role`);',
        '--> statement-breakpoint',
        'CREATE INDEX `idx_control_cabinet_items_component` ON `control_cabinet_items` (`component_id`);',
        '--> statement-breakpoint',
    ]
    for component in components:
        lines.append(
            "INSERT INTO control_components "
            "(id,category,manufacturer,article,name,component_type,attributes_json,current_price_microunits,currency,source_file,source_sheet,source_row) VALUES "
            f"({sql_string(component['id'])},{sql_string(component['category'])},{sql_string(component['manufacturer'])},"
            f"{sql_string(component['article'])},{sql_string(component['name'])},{sql_string(component['componentType'])},"
            f"{sql_string(json.dumps(component['attributes'], ensure_ascii=False, separators=(',', ':')))},"
            f"{micros(component['currentPrice'])},{sql_string(component['currency'])},{sql_string(component['sourceFile'])},"
            f"{sql_string(component['sourceSheet'])},{component['sourceRow']});"
        )
        lines.append("--> statement-breakpoint")
    for cabinet in cabinets:
        lines.append(
            "INSERT INTO control_cabinets "
            "(id,configuration_key,station_type,name,pump_count,pump_power_kw,breaker_current_a,incoming_switch_current_a,"
            "contactor_count,vfd_count,vfd_power_kw,enclosure_dimensions,assembly_kit_type,labor_hours,labor_rate,"
            "cached_total_microunits,source) VALUES "
            f"({sql_string(cabinet['id'])},{sql_string(cabinet['configurationKey'])},{sql_string(cabinet['stationType'])},"
            f"{sql_string(cabinet['name'])},{cabinet['pumpCount']},{cabinet['pumpPowerKw']},{cabinet['breakerCurrentA']},"
            f"{cabinet['incomingSwitchCurrentA']},{cabinet['contactorCount']},{cabinet['vfdCount']},{cabinet['vfdPowerKw']},"
            f"{sql_string(cabinet['enclosureDimensions'])},{sql_string(cabinet['assemblyKitType'])},{cabinet['laborHours']},"
            f"{cabinet['laborRate']},{micros(cabinet['cachedTotal'])},{sql_string(cabinet['source'])});"
        )
        lines.append("--> statement-breakpoint")
        for item in cabinet["items"]:
            lines.append(
                "INSERT INTO control_cabinet_items (cabinet_id,component_id,role,quantity,sort_order) VALUES "
                f"({sql_string(cabinet['id'])},{sql_string(item['componentId'])},{sql_string(item['role'])},"
                f"{item['quantity']},{item['sortOrder']});"
            )
            lines.append("--> statement-breakpoint")
    SQL_OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    components = read_components()
    cabinets = read_cabinets(components)
    payload = {
        "schemaVersion": 1,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "source": str(SOURCE_ROOT.relative_to(ROOT)).replace("\\", "/"),
        "statistics": {
            "sourceFiles": len(price_files()),
            "components": len(components),
            "pricedComponents": sum(item["currentPrice"] is not None for item in components),
            "readyCabinets": len(cabinets),
            "cabinetItems": sum(len(item["items"]) for item in cabinets),
        },
        "components": components,
        "cabinets": cabinets,
    }
    PUBLIC_OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    write_sql(components, cabinets)
    print(json.dumps(payload["statistics"], ensure_ascii=False))
    for cabinet in cabinets:
        difference = round(cabinet["currentTotal"] - cabinet["cachedTotal"], 2)
        print(f"{cabinet['id']}: cached={cabinet['cachedTotal']:.2f}, current={cabinet['currentTotal']:.2f}, delta={difference:+.2f}")


if __name__ == "__main__":
    main()
