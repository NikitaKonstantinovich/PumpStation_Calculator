from __future__ import annotations

import json
import math
import re
import sys
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "Equipment" / "Калькулятор по ОБВЯЗКЕ.xlsm"
OUTPUT = ROOT / "frontend" / "public" / "binding-components.json"

CATALOG_SHEETS = [
    "Обвязка",
    "Сгоны резьба",
    "Трубы",
    "Металл",
    "Фланцы",
    "Воротники",
    "Отводы",
    "Переходы",
    "Тройники",
    "Заглушки",
    "Мембранный бак",
    "Метизы",
    "Прокладки",
    "Опоры",
    "КИПиА",
    "Кронштейны",
    "Рама",
    "Работы",
]

RULE_SHEETS = [
    "Исходные данные",
    "СПЕЦИФИКАЦИЯ СТАНЦИЯ",
    "Спецификация пожарная станция",
    "Пожарная станция",
    *CATALOG_SHEETS,
    "Полезная инфа",
    "Для зависимостей",
]

HEADER_WORDS = re.compile(
    r"(?:^|\b)(dn|pn|ду|ру|цена|масса|вес|наименование|диаметр|длина|толщина|"
    r"количество|давление|материал|исполнение|тип|размер|резьба|мощность|ссылка|"
    r"артикул|обозначение|единица|норма)(?:\b|,)",
    re.IGNORECASE,
)
PRICE_WORDS = re.compile(r"цен|стоим", re.IGNORECASE)
FORMULA_REF = re.compile(
    r"(?:(?:'([^']+)'|([\wА-Яа-яЁё .()№-]+))!)?\$?([A-Z]{1,3})\$?(\d+)",
    re.UNICODE,
)

OBVYAZKA_TABLES = {
    1: ("затворы", "Затворы", "Затвор поворотный дисковый 017W · Синий PN10/16"),
    2: ("затворы", "Затворы", "Затвор поворотный дисковый 017W · Красный PN16"),
    3: ("затворы", "Затворы", "Затвор поворотный дисковый 017W · Красный PN25"),
    4: ("задвижки", "Задвижки", "Задвижка клиновая фланцевая 47GV · Синяя"),
    5: ("задвижки", "Задвижки", "Задвижка клиновая фланцевая 47GV · Красная"),
    6: ("обратные-клапаны", "Обратные клапаны", "Клапан обратный двустворчатый межфланцевый 010С.Y · Синий"),
    7: ("обратные-клапаны", "Обратные клапаны", "Клапан обратный двустворчатый межфланцевый 010С.Y · Красный"),
    8: ("обратные-клапаны", "Обратные клапаны", "Клапан обратный шаровой фланцевый 012F"),
    9: ("компенсаторы", "Компенсаторы", "Компенсатор резиновый фланцевый KMS · EPDM"),
    10: ("электроприводы", "Электроприводы", "Электропривод многооборотный · IP67, 380 В"),
    11: ("задвижки", "Задвижки", "Задвижка шиберная межфланцевая K21GV · Синяя"),
    12: ("обратные-клапаны", "Обратные клапаны", "Обратный клапан с латунным сердечником"),
    13: ("шаровые-краны", "Шаровые краны", "Кран шаровой латунный ВР · ручка-рычаг"),
    14: ("шаровые-краны", "Шаровые краны", "Кран шаровой латунный ВР-ВР · дренаж и воздухоотводчик"),
    15: ("шаровые-краны", "Шаровые краны", "Кран шаровой латунный ВР · ручка-бабочка"),
    16: ("резьбовые-фитинги", "Резьбовые фитинги", "Ниппель латунный НР/НР · LD Pride"),
    17: ("резьбовые-фитинги", "Резьбовые фитинги", "Футорка латунная НР-ВР · LD Pride"),
    18: ("резьбовые-фитинги", "Резьбовые фитинги", "Переходник латунный ВР-НР · MVI"),
    19: ("пожарная-арматура", "Пожарная арматура", "Головка пожарная муфтовая ГМ ПМ УХЛ1 · Цветлит"),
    20: ("пожарная-арматура", "Пожарная арматура", "Головка-заглушка ГЗ А-П"),
}

OBVYAZKA_COLUMN_OVERRIDES = {
    6: [["DN"], ["Чугун", "Цена, ₽"], ["Нержавеющая сталь", "Цена, ₽"], ["Вес, кг"]],
    7: [["DN"], ["Чугун", "Цена, ₽"], ["Нержавеющая сталь", "Цена, ₽"], ["Вес, кг"]],
    8: [["DN"], ["Тип 012F.Z", "Чугун", "Цена, ₽"], ["Тип 012F.M", "Чугун", "Цена, ₽"]],
    9: [["DN"], ["PN10", "Цена, ₽"], ["PN16", "Цена, ₽"], ["PN25", "Цена, ₽"], ["PN10", "Вес, кг"], ["PN16", "Вес, кг"], ["PN25", "Вес, кг"]],
    10: [["Тип привода"], ["IP67 · 380 В", "Цена, ₽"], ["Вес, кг"]],
    12: [["DN"], ["Латунь", "Цена, ₽"], ["Ссылка"]],
}

OBVYAZKA_EXTRA_HEADER_ROWS = {
    6: {59, 60},
    7: {59, 60},
    8: {85, 86},
    9: {110, 111},
    10: {135},
    12: {182},
}


def clean_scalar(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return round(value, 10)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.replace("\u00a0", " ").replace("\r\n", "\n").strip()
    return str(value)


def numeric_value(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return round(float(value), 10)
    if not isinstance(value, str):
        return None
    text = value.strip().replace("\u00a0", " ").replace("�", " ")
    if not text or text.startswith("#"):
        return None
    if not re.fullmatch(r"[-+]?\s*\d[\d\s]*(?:[.,]\d+)?(?:\s*[₽рР])?", text):
        return None
    normalized = re.sub(r"[\s₽рР]", "", text).replace(",", ".")
    try:
        return round(float(normalized), 10)
    except ValueError:
        return None


def slug(value: str) -> str:
    value = value.casefold().replace("ё", "е")
    value = re.sub(r"[^a-zа-я0-9]+", "-", value, flags=re.IGNORECASE).strip("-")
    return value or "item"


def workbook_xml_maps(source: Path) -> tuple[
    dict[str, str],
    dict[str, dict[str, dict[str, Any]]],
    dict[str, list[dict[str, Any]]],
    list[dict[str, str]],
]:
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(source) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.attrib["Id"]: rel.attrib["Target"]
            for rel in rels_root.findall(f"{{{ns_pkg}}}Relationship")
        }
        sheet_paths: dict[str, str] = {}
        for sheet in workbook_root.findall(f".//{{{ns_main}}}sheet"):
            name = sheet.attrib["name"].strip()
            target = rel_targets[sheet.attrib[f"{{{ns_rel}}}id"]]
            sheet_paths[name] = "xl/" + target.lstrip("/")

        formulas: dict[str, dict[str, dict[str, Any]]] = {}
        extended_validations: dict[str, list[dict[str, Any]]] = {}
        for sheet_name, xml_path in sheet_paths.items():
            root = ET.fromstring(archive.read(xml_path))
            current: dict[str, dict[str, Any]] = {}
            for cell in root.findall(f".//{{{ns_main}}}c"):
                formula = cell.find(f"{{{ns_main}}}f")
                if formula is None:
                    continue
                cached = cell.find(f"{{{ns_main}}}v")
                current[cell.attrib["r"]] = {
                    "formula": clean_scalar(formula.text or ""),
                    "cachedRaw": clean_scalar(cached.text if cached is not None else None),
                }
            formulas[sheet_name] = current
            extended_validations[sheet_name] = []
            x14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
            xm = "http://schemas.microsoft.com/office/excel/2006/main"
            for validation in root.findall(f".//{{{x14}}}dataValidation"):
                formula_node = validation.find(f".//{{{xm}}}f")
                range_node = validation.find(f".//{{{xm}}}sqref")
                extended_validations[sheet_name].append({
                    "range": clean_scalar(range_node.text if range_node is not None else None),
                    "type": validation.attrib.get("type"),
                    "operator": validation.attrib.get("operator"),
                    "formula1": clean_scalar(formula_node.text if formula_node is not None else None),
                    "formula2": None,
                    "allowBlank": validation.attrib.get("allowBlank") == "1",
                    "extension": "x14",
                })

        external_links: list[dict[str, str]] = []
        for name in sorted(n for n in archive.namelist() if re.fullmatch(r"xl/externalLinks/_rels/externalLink\d+\.xml\.rels", n)):
            root = ET.fromstring(archive.read(name))
            for rel in root.findall(f"{{{ns_pkg}}}Relationship"):
                external_links.append({
                    "id": rel.attrib.get("Id", ""),
                    "target": rel.attrib.get("Target", ""),
                    "targetMode": rel.attrib.get("TargetMode", ""),
                })
    return sheet_paths, formulas, extended_validations, external_links


def merged_anchor_map(worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in worksheet.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                anchors[(row, col)] = (merged.min_row, merged.min_col)
    return anchors


def expanded_value(worksheet, anchors: dict[tuple[int, int], tuple[int, int]], row: int, col: int) -> Any:
    anchor = anchors.get((row, col), (row, col))
    return clean_scalar(worksheet.cell(*anchor).value)


def row_groups(worksheet, anchors: dict[tuple[int, int], tuple[int, int]]) -> list[tuple[int, int]]:
    used_rows = []
    for row in range(1, worksheet.max_row + 1):
        if any(expanded_value(worksheet, anchors, row, col) not in (None, "") for col in range(1, worksheet.max_column + 1)):
            used_rows.append(row)
    if not used_rows:
        return []
    groups: list[tuple[int, int]] = []
    start = previous = used_rows[0]
    for row in used_rows[1:]:
        if row > previous + 1:
            groups.append((start, previous))
            start = row
        previous = row
    groups.append((start, previous))
    return groups


def column_groups(worksheet, anchors: dict[tuple[int, int], tuple[int, int]], row_start: int, row_end: int) -> list[tuple[int, int]]:
    used_cols = []
    for col in range(1, worksheet.max_column + 1):
        if any(expanded_value(worksheet, anchors, row, col) not in (None, "") for row in range(row_start, row_end + 1)):
            used_cols.append(col)
    if not used_cols:
        return []
    groups: list[tuple[int, int]] = []
    start = previous = used_cols[0]
    for col in used_cols[1:]:
        if col > previous + 1:
            groups.append((start, previous))
            start = col
        previous = col
    groups.append((start, previous))
    return groups


def row_values(worksheet, row: int, col_start: int, col_end: int) -> list[Any]:
    return [clean_scalar(worksheet.cell(row, col).value) for col in range(col_start, col_end + 1)]


def detect_header_end(matrix: list[list[Any]]) -> int:
    candidates: list[int] = []
    for index, row in enumerate(matrix[:10]):
        text = " | ".join(str(value) for value in row if value not in (None, ""))
        if HEADER_WORDS.search(text):
            candidates.append(index)
    if candidates:
        header_end = min(candidates)
        for index in range(header_end + 1, min(len(matrix), header_end + 4)):
            text = " | ".join(str(value) for value in matrix[index] if value not in (None, ""))
            if HEADER_WORDS.search(text) and not re.search(r"\d", text):
                header_end = index
                continue
            break
        return header_end
    return 0 if len(matrix) > 1 else -1


def header_paths(expanded_matrix: list[list[Any]], header_end: int, width: int) -> list[list[str]]:
    paths: list[list[str]] = []
    for col in range(width):
        values = []
        for row in range(header_end + 1):
            value = expanded_matrix[row][col]
            if value in (None, ""):
                continue
            text = str(value).strip()
            if text.casefold() == "вернуться" or text in values:
                continue
            values.append(text)
        paths.append(values)
    return paths


def formula_dependencies(formula: str, current_sheet: str) -> list[dict[str, str]]:
    dependencies = []
    seen = set()
    for match in FORMULA_REF.finditer(formula):
        sheet = (match.group(1) or match.group(2) or current_sheet).strip()
        address = f"{match.group(3)}{match.group(4)}"
        key = (sheet, address)
        if key in seen:
            continue
        seen.add(key)
        dependencies.append({"sheet": sheet, "address": address})
    return dependencies


def source_cell(worksheet_formula, worksheet_values, formula_map, row: int, col: int, header_path: list[str]) -> dict[str, Any]:
    address = f"{get_column_letter(col)}{row}"
    formula_info = formula_map.get(address)
    raw_formula_value = None if isinstance(worksheet_formula.cell(row, col), MergedCell) else clean_scalar(worksheet_formula.cell(row, col).value)
    formula = formula_info["formula"] if formula_info else None
    if formula_info is not None and not formula and isinstance(raw_formula_value, str) and raw_formula_value.startswith("="):
        formula = raw_formula_value[1:]
    cached = clean_scalar(worksheet_values.cell(row, col).value)
    value = cached if formula is not None else raw_formula_value
    numeric = numeric_value(value)
    cell = worksheet_formula.cell(row, col)
    result: dict[str, Any] = {
        "column": get_column_letter(col),
        "headerPath": header_path,
        "value": value,
    }
    if numeric is not None and not isinstance(value, (int, float)):
        result["numericValue"] = numeric
    if formula is not None:
        result["formula"] = formula
        result["dependencies"] = formula_dependencies(formula, worksheet_formula.title)
    if cell.number_format and cell.number_format != "General":
        result["numberFormat"] = cell.number_format
    if cell.hyperlink:
        result["sourceUrl"] = cell.hyperlink.target
    if cell.comment:
        result["comment"] = clean_scalar(cell.comment.text)
    return result


def parse_tables(worksheet_formula, worksheet_values, formula_map, sheet_id: str) -> list[dict[str, Any]]:
    anchors = merged_anchor_map(worksheet_formula)
    tables = []
    table_number = 0
    for row_start, row_end in row_groups(worksheet_formula, anchors):
        for col_start, col_end in column_groups(worksheet_formula, anchors, row_start, row_end):
            matrix = [row_values(worksheet_values, row, col_start, col_end) for row in range(row_start, row_end + 1)]
            nonempty = sum(value not in (None, "") for row in matrix for value in row)
            if nonempty < 3:
                continue
            expanded_matrix = [
                [expanded_value(worksheet_formula, anchors, row, col) for col in range(col_start, col_end + 1)]
                for row in range(row_start, row_end + 1)
            ]
            header_end = detect_header_end(matrix)
            paths = header_paths(expanded_matrix, header_end, col_end - col_start + 1)
            title_parts = []
            for row in matrix[: max(1, header_end + 1)]:
                for value in row:
                    if not isinstance(value, str) or not value.strip() or value.casefold() == "вернуться":
                        continue
                    if HEADER_WORDS.search(value):
                        continue
                    if value not in title_parts:
                        title_parts.append(value)
            title = " · ".join(title_parts[:6]) or worksheet_formula.title
            records = []
            for offset, row_values_cached in enumerate(matrix[header_end + 1 :], start=header_end + 1):
                source_row = row_start + offset
                if all(value in (None, "") for value in row_values_cached):
                    continue
                fields = []
                for local_col, value in enumerate(row_values_cached):
                    if value in (None, ""):
                        continue
                    actual_col = col_start + local_col
                    fields.append(source_cell(worksheet_formula, worksheet_values, formula_map, source_row, actual_col, paths[local_col]))
                if not fields:
                    continue
                record_id = f"{sheet_id}-{table_number + 1:03d}-r{source_row:04d}"
                prices = []
                source_urls = []
                for field in fields:
                    header_text = " / ".join(field["headerPath"])
                    amount = numeric_value(field["value"])
                    if amount is not None and amount > 0 and PRICE_WORDS.search(header_text):
                        prices.append({
                            "label": header_text or "Цена",
                            "amount": amount,
                            "currency": "RUB",
                            "sourceColumn": field["column"],
                        })
                    if field.get("sourceUrl") and field["sourceUrl"] not in source_urls:
                        source_urls.append(field["sourceUrl"])
                records.append({
                    "id": record_id,
                    "sourceRow": source_row,
                    "fields": fields,
                    "prices": prices,
                    "sourceUrls": source_urls,
                })
            table_number += 1
            table_urls = list(dict.fromkeys(url for record in records for url in record["sourceUrls"]))
            tables.append({
                "id": f"{sheet_id}-{table_number:03d}",
                "title": title,
                "sourceRange": f"{get_column_letter(col_start)}{row_start}:{get_column_letter(col_end)}{row_end}",
                "headerRows": list(range(row_start, row_start + header_end + 1)),
                "columns": [
                    {"column": get_column_letter(col_start + index), "headerPath": path}
                    for index, path in enumerate(paths)
                ],
                "sourceUrls": table_urls,
                "records": records,
            })
    return tables


def rebuild_prices(fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
    prices = []
    for field in fields:
        label = " / ".join(field["headerPath"])
        amount = numeric_value(field["value"])
        if amount is not None and amount > 0 and PRICE_WORDS.search(label):
            prices.append({
                "label": label or "Цена",
                "amount": amount,
                "currency": "RUB",
                "sourceColumn": field["column"],
            })
    return prices


def normalize_obvyazka(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Turn the mixed source sheet into clean, product-oriented catalogues."""
    grouped: dict[str, dict[str, Any]] = {}
    counters: Counter[str] = Counter()
    for source_number, table in enumerate(tables, start=1):
        category_id, category_name, title = OBVYAZKA_TABLES[source_number]
        counters[category_id] += 1
        new_table_id = f"{category_id}-{counters[category_id]:03d}"
        table["id"] = new_table_id
        table["title"] = title

        override_paths = OBVYAZKA_COLUMN_OVERRIDES.get(source_number)
        if override_paths:
            for column, header_path in zip(table["columns"], override_paths):
                column["headerPath"] = header_path
            paths_by_column = {column["column"]: column["headerPath"] for column in table["columns"]}
            for record in table["records"]:
                for field in record["fields"]:
                    if field["column"] in paths_by_column:
                        field["headerPath"] = paths_by_column[field["column"]]

        extra_headers = OBVYAZKA_EXTRA_HEADER_ROWS.get(source_number, set())
        table["headerRows"] = sorted({*table["headerRows"], *extra_headers})
        records = []
        seen_products = set()
        for record in table["records"]:
            if record["sourceRow"] in extra_headers:
                continue
            record["prices"] = rebuild_prices(record["fields"])
            # Rows with only a size and placeholders are incomplete source templates,
            # not purchasable components. They made the viewer look like it contained
            # products named "-" and are intentionally omitted.
            if not record["prices"]:
                continue
            identity = tuple(
                str(field["value"]).strip().casefold()
                for field in record["fields"]
                if field["value"] not in (None, "", "-")
            )
            # The drive table repeats the same component for several valve sizes.
            # In a component catalogue a drive model is a single product.
            if source_number == 10:
                identity = identity[:1]
            if identity in seen_products:
                continue
            seen_products.add(identity)
            record["id"] = f"{new_table_id}-r{record['sourceRow']:04d}"
            records.append(record)
        table["records"] = records

        if not records:
            continue
        category = grouped.setdefault(category_id, {
            "id": category_id,
            "name": category_name,
            "sourceSheet": "Обвязка",
            "tables": [],
        })
        category["tables"].append(table)
    return list(grouped.values())


def extract_validations(worksheet) -> list[dict[str, Any]]:
    validations = []
    for validation in worksheet.data_validations.dataValidation:
        validations.append({
            "range": str(validation.sqref),
            "type": validation.type,
            "operator": validation.operator,
            "formula1": clean_scalar(validation.formula1),
            "formula2": clean_scalar(validation.formula2),
            "allowBlank": validation.allow_blank,
        })
    return validations


def defined_names(workbook) -> list[dict[str, Any]]:
    result = []
    for name, definition in workbook.defined_names.items():
        result.append({
            "name": name,
            "value": clean_scalar(definition.attr_text),
            "localSheetId": definition.localSheetId,
            "hidden": bool(definition.hidden),
        })
    return result


def collector_reference(workbook) -> dict:
    sheet = next(sheet for sheet in workbook if sheet.title.strip() == "Полезная инфа")
    # Explicit source ranges: two-flange bolt length, and material pipe OD.
    # Do not use the unrelated valve/check-valve bolt-length tables.
    if "Длина болтов" not in str(sheet["P49"].value):
        raise ValueError("Изменилась структура таблицы длин болтов: Полезная инфа!P49")
    bolts = []
    for row in range(51, 71):
        dn = int(str(sheet[f"B{row}"].value).removeprefix("DN"))
        for column, pn in [("P", 10), ("Q", 16), ("R", 25)]:
            length = numeric_value(sheet[f"{column}{row}"].value)
            if length and length > 0:
                bolts.append({"dn": dn, "pn": pn, "length": length, "source": f"Полезная инфа!{column}{row} (два фланца)"})
    diameters = []
    for material, rows in [("st20", range(4, 24)), ("aisi304", range(28, 48))]:
        for row in rows:
            dn = int(str(sheet[f"B{row}"].value).removeprefix("DN"))
            # Sgon DN15/DN20 uses the matching pipe outside diameter from this table.
            if dn in (15, 20):
                diameters.append({"dn": dn, "material": material, "outerDiameter": sheet[f"C{row}"].value, "source": f"Полезная инфа!C{row}"})
    return {"bolts": bolts, "nippleDiameters": diameters}


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    sheet_paths, xml_formulas, extended_validations, external_links = workbook_xml_maps(SOURCE)
    workbook_formula = load_workbook(SOURCE, data_only=False, keep_vba=True, keep_links=True)
    workbook_values = load_workbook(SOURCE, data_only=True, keep_vba=True, keep_links=True)
    canonical_sheet_names = list(sheet_paths)
    if len(canonical_sheet_names) != len(workbook_formula.worksheets):
        raise ValueError("Количество листов в XML и книге не совпадает")
    for index, sheet_name in enumerate(canonical_sheet_names):
        workbook_formula.worksheets[index].title = sheet_name
        workbook_values.worksheets[index].title = sheet_name

    catalogs = []
    flat_items = []
    for sheet_name in CATALOG_SHEETS:
        sheet_id = slug(sheet_name)
        tables = parse_tables(
            workbook_formula[sheet_name],
            workbook_values[sheet_name],
            xml_formulas.get(sheet_name, {}),
            sheet_id,
        )
        source_catalogs = normalize_obvyazka(tables) if sheet_name == "Обвязка" else [{
            "id": sheet_id,
            "name": sheet_name,
            "sourceSheet": sheet_name,
            "tables": tables,
        }]
        for source_catalog in source_catalogs:
            catalog_id = source_catalog["id"]
            catalog_tables = source_catalog["tables"]
            catalog = {
                "id": catalog_id,
                "name": source_catalog["name"],
                "sourceSheet": source_catalog["sourceSheet"],
                "tables": [
                    {
                        **{key: value for key, value in table.items() if key != "records"},
                        "recordIds": [record["id"] for record in table["records"]],
                    }
                    for table in catalog_tables
                ],
            }
            catalogs.append(catalog)
            for table in catalog_tables:
                for record in table["records"]:
                    flat_items.append({
                        "id": record["id"],
                        "catalogId": catalog_id,
                        "tableId": table["id"],
                        "family": table["title"],
                        "sourceSheet": sheet_name,
                        "sourceRow": record["sourceRow"],
                        "fields": record["fields"],
                        "prices": record["prices"],
                        "sourceUrls": list(dict.fromkeys([*table["sourceUrls"], *record["sourceUrls"]])),
                    })

    formula_rules = []
    for sheet_name in RULE_SHEETS:
        value_sheet = workbook_values[sheet_name]
        for address, formula_info in xml_formulas.get(sheet_name, {}).items():
            cached_value = clean_scalar(value_sheet[address].value)
            formula = formula_info["formula"]
            if not formula:
                openpyxl_formula = clean_scalar(workbook_formula[sheet_name][address].value)
                formula = openpyxl_formula[1:] if isinstance(openpyxl_formula, str) and openpyxl_formula.startswith("=") else ""
            formula_rules.append({
                "id": f"{slug(sheet_name)}-{address.casefold()}",
                "sheet": sheet_name,
                "address": address,
                "formula": formula,
                "cachedValue": cached_value,
                "dependencies": formula_dependencies(formula, sheet_name),
            })

    dependency_tables = parse_tables(
        workbook_formula["Для зависимостей"],
        workbook_values["Для зависимостей"],
        xml_formulas.get("Для зависимостей", {}),
        "dependencies",
    )
    input_tables = parse_tables(
        workbook_formula["Исходные данные"],
        workbook_values["Исходные данные"],
        xml_formulas.get("Исходные данные", {}),
        "inputs",
    )
    validations = [
        *extract_validations(workbook_formula["Исходные данные"]),
        *extended_validations.get("Исходные данные", []),
    ]

    priced_items = [item for item in flat_items if item["prices"]]
    price_entries = sum(len(item["prices"]) for item in flat_items)
    formula_error_count = sum(
        isinstance(rule["cachedValue"], str) and rule["cachedValue"].startswith("#")
        for rule in formula_rules
    )
    data = {
        "schemaVersion": 1,
        "source": {
            "file": "Equipment/Калькулятор по ОБВЯЗКЕ.xlsm",
            "sha256": source_hash,
            "sizeBytes": SOURCE.stat().st_size,
            "modifiedAt": datetime.fromtimestamp(SOURCE.stat().st_mtime, timezone.utc).isoformat(),
            "importedAt": datetime.now(timezone.utc).isoformat(),
            "currency": "RUB",
            "externalLinks": external_links,
        },
        "statistics": {
            "catalogs": len(catalogs),
            "tables": sum(len(catalog["tables"]) for catalog in catalogs),
            "componentRows": len(flat_items),
            "pricedComponentRows": len(priced_items),
            "priceEntries": price_entries,
            "formulaRules": len(formula_rules),
            "formulaErrorsInSavedValues": formula_error_count,
            "dataValidations": len(validations),
            "definedNames": len(workbook_formula.defined_names),
        },
        "catalogs": catalogs,
        "items": flat_items,
        "collectorReference": collector_reference(workbook_values),
        "relations": {
            "dependencyTables": dependency_tables,
            "inputTables": input_tables,
            "dataValidations": validations,
            "definedNames": defined_names(workbook_formula),
            "formulas": formula_rules,
        },
        "quality": {
            "notes": [
                "Цена хранится вместе с валютой, условием варианта и координатой исходной ячейки.",
                "Для формул сохранены выражение, последнее рассчитанное значение и список прямых ссылок на ячейки.",
                "Внешние ссылки не обновлялись: импортированы значения, сохранённые в исходной книге.",
            ],
            "catalogRecordCounts": dict(Counter(item["catalogId"] for item in flat_items)),
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), **data["statistics"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Import failed: {error}", file=sys.stderr)
        raise
